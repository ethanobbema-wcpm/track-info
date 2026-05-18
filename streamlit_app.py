from __future__ import annotations

import base64
import re
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


APP_TITLE = "Track Info"
INSTRUMENT_TAXONOMY_PATH = Path(__file__).with_name("taxonomy_instruments.xlsx")
VOCAL_SUBLIST_PATH = Path(__file__).with_name("vocal_sublist.xlsx")
KEY_SELECTION_PATH = Path(__file__).with_name("key_selection.xlsx")
LOGO_PATH = Path(__file__).with_name("wcpm_logo.png")
TITLE_FONT_PATH = Path(__file__).with_name("plaak_title.ttf")
BODY_FONT_PATH = Path(__file__).with_name("bw_gradual_light.otf")
BODY_FONT_MEDIUM_PATH = Path(__file__).with_name("bw_gradual_medium.otf")
WORDPROCESSINGML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

INTRO_TEXT = (
    "Please, ensure that the provided information is meticulously verified. "
    "Always specify keys as major or minor, such as F and Fm, and the BPM "
    "should accurately represent the pulse's overall feel, which typically "
    "corresponds to the quarter note value rather than the eighth note value. "
    "There might be a few occasions where the half note value would make more "
    "sense. If you are arranging a Public Domain work, please, indicate the "
    "EXACT original title between parentheses. Please, be precise when listing "
    "instruments. "
)
INTRO_TEXT_EMPHASIS = (
    "Export the zip at the bottom of the page. Please do not edit the exported files "
    "and instead use the FIX button to make any changes before exporting again."
)
HELP_TOOLTIP_ITEMS = [
    "Exporting is only available after all required fields are completed.",
    "Expand the Required Fields dropdown at the bottom of the page to see what remaining fields must be completed.",
    'The "Reset" button clears the page.',
    'Use the "Need to fix Track Info?" button to upload a previously exported zip for editing.',
]

COMPOSER_FIELDS = {
    "composer": "Composer:",
    "pro_affiliation": "PRO Affiliation:",
    "cae_ipi": "CAE/IPI",
    "split": "Splits:",
    "spotify_link": "DSP Links - Spotify:",
    "apple_music_link": "Apple Music:",
    "artist_name": "Artist Name (Optional):",
}
COMPOSER_TEXT_FIELDS = (
    "composer",
    "pro_affiliation",
    "spotify_link",
    "apple_music_link",
    "artist_name",
)

BASE_HEADERS = [
    "Album Name",
    "WCPM Producer",
    "Track Number",
    "Track Title",
    "BPM",
    "Key",
    "Meter",
    "Instrumentation",
    "Vocal Sub-list",
    "Featured Instrument",
]

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
TRACK_HEADER_COLORS = [
    "#FFADAD",
    "#FFD6A5",
    "#FDFFB6",
    "#CAFFBF",
    "#9BF6FF",
    "#A0C4FF",
    "#BDB2FF",
    "#FFC6FF",
]


def build_track_header_css(max_track_count: int = 200) -> str:
    css_rules: list[str] = []
    for track_number in range(1, max_track_count + 1):
        color = TRACK_HEADER_COLORS[(track_number - 1) % len(TRACK_HEADER_COLORS)]
        css_rules.append(
            f"""
            .st-key-track-panel-{track_number} div[data-testid="stExpander"] details summary {{
                background: {color};
                border-radius: 8px 8px 0 0;
                color: #1f2937;
            }}
            .st-key-track-panel-{track_number} div[data-testid="stExpander"] details summary:hover {{
                background: color-mix(in srgb, {color} 92%, #000000);
                color: #111827;
            }}
            .st-key-track-panel-{track_number} div[data-testid="stExpander"] details[open] summary {{
                background: color-mix(in srgb, {color} 88%, #000000);
                color: #111827;
            }}
            .st-key-track-panel-{track_number} div[data-testid="stExpander"] details summary * {{
                color: inherit !important;
                opacity: 1 !important;
            }}
            """
        )

    return "\n".join(css_rules)


def configure_page() -> None:
    track_header_css = build_track_header_css()
    title_font_css = ""
    body_font_css = ""
    if TITLE_FONT_PATH.exists():
        title_font_data_uri = load_binary_asset_data_uri(
            str(TITLE_FONT_PATH),
            TITLE_FONT_PATH.stat().st_mtime_ns,
        )
        title_font_css = f"""
            @font-face {{
                font-family: "Plaak Title";
                src: url("{title_font_data_uri}") format("truetype");
                font-style: normal;
                font-weight: 800;
                font-display: swap;
            }}
            div[data-testid="stHeadingWithActionElements"] h1 {{
                font-family: "Plaak Title", var(--font, sans-serif);
                letter-spacing: 0;
            }}
        """
    if BODY_FONT_PATH.exists():
        body_font_data_uri = load_binary_asset_data_uri(
            str(BODY_FONT_PATH),
            BODY_FONT_PATH.stat().st_mtime_ns,
        )
        body_font_medium_css = ""
        if BODY_FONT_MEDIUM_PATH.exists():
            body_font_medium_data_uri = load_binary_asset_data_uri(
                str(BODY_FONT_MEDIUM_PATH),
                BODY_FONT_MEDIUM_PATH.stat().st_mtime_ns,
            )
            body_font_medium_css = f"""
                @font-face {{
                    font-family: "Bw Gradual Medium";
                    src: url("{body_font_medium_data_uri}") format("opentype");
                    font-style: normal;
                    font-weight: 500;
                    font-display: swap;
                }}
                .intro-copy strong,
                .intro-copy .intro-emphasis {{
                    font-family: "Bw Gradual Medium", "Bw Gradual", var(--font, sans-serif);
                    font-weight: 500;
                }}
            """
        body_font_css = f"""
            @font-face {{
                font-family: "Bw Gradual";
                src: url("{body_font_data_uri}") format("opentype");
                font-style: normal;
                font-weight: 300;
                font-display: swap;
            }}
            .stApp {{
                font-family: "Bw Gradual", var(--font, sans-serif);
            }}
            .stApp button,
            .stApp input,
            .stApp textarea,
            .stApp select,
            .stApp label,
            .stApp [data-baseweb="select"],
            .stApp [data-baseweb="input"],
            .stApp [data-baseweb="textarea"],
            .stApp [data-testid="stMarkdownContainer"],
            .stApp [data-testid="stText"],
            .stApp [data-testid="stWidgetLabel"],
            .stApp [data-testid="stExpander"] summary {{
                font-family: inherit;
            }}
            {body_font_medium_css}
        """
    base_css = """
        <style>
            .intro-copy {
                color: var(--text-color, inherit);
                font-size: 0.95rem;
                line-height: 1.55;
                margin-top: -0.35rem;
                max-width: 100%;
                opacity: 0.92;
            }
            .app-logo-wrap {
                display: flex;
                justify-content: center;
                margin: 0.15rem 0 0.85rem;
            }
            .app-logo {
                display: block;
                height: auto;
                max-width: 100%;
                width: 336px;
            }
            .page-help-row {
                display: flex;
                justify-content: flex-end;
                margin: 0 0 0.35rem;
            }
            .page-help {
                position: relative;
                display: inline-flex;
                justify-content: center;
                align-items: center;
            }
            .page-help-icon {
                align-items: center;
                background: var(--secondary-background-color, rgba(255, 255, 255, 0.06));
                border: 1px solid rgba(148, 163, 184, 0.45);
                border-radius: 999px;
                color: var(--text-color, inherit);
                cursor: default;
                display: inline-flex;
                font-family: "Bw Gradual Medium", "Bw Gradual", var(--font, sans-serif);
                font-size: 0.95rem;
                height: 1.9rem;
                justify-content: center;
                line-height: 1;
                width: 1.9rem;
            }
            .page-help-tooltip {
                background: #ffffff;
                border: 1px solid rgba(15, 23, 42, 0.12);
                border-radius: 8px;
                box-shadow: 0 12px 30px rgba(15, 23, 42, 0.28);
                color: #111827;
                opacity: 0;
                padding: 0.8rem 1rem;
                pointer-events: none;
                position: absolute;
                right: 0;
                top: calc(100% + 0.45rem);
                transform: translateY(-0.2rem);
                transition: opacity 140ms ease, transform 140ms ease;
                visibility: hidden;
                width: min(420px, 72vw);
                z-index: 50;
            }
            .page-help:hover .page-help-tooltip,
            .page-help:focus-within .page-help-tooltip {
                opacity: 1;
                transform: translateY(0);
                visibility: visible;
            }
            .page-help-tooltip ul {
                margin: 0;
                padding-left: 1.1rem;
            }
            .page-help-tooltip li {
                color: #111827;
                line-height: 1.45;
                margin: 0 0 0.45rem;
            }
            .page-help-tooltip li:last-child {
                margin-bottom: 0;
            }
            div[data-testid="stExpander"] {
                border-radius: 8px;
            }
            div[data-testid="stButton"] > button,
            div[data-testid="stDownloadButton"] > button {
                border-radius: 8px;
                min-height: 2.6rem;
            }
            div[data-testid="stButton"] > button[kind="tertiary"] {
                background: transparent;
                border: none;
                color: var(--text-color, inherit);
                justify-content: flex-start;
                min-height: 1.1rem;
                opacity: 0.82;
                padding: 0;
                text-decoration: underline;
                white-space: nowrap;
            }
            div[data-testid="stButton"] > button[kind="tertiary"]:hover {
                background: transparent;
                color: var(--text-color, inherit);
                opacity: 1;
            }
            .st-key-track-info-import-toggle div[data-testid="stButton"] > button {
                background: #facc15;
                border: 1px solid #eab308;
                color: #111827;
                font-weight: 600;
            }
            .st-key-track-info-import-toggle div[data-testid="stButton"] > button:hover {
                background: #fbbf24;
                border-color: #f59e0b;
                color: #111827;
            }
            .st-key-track-info-reset div[data-testid="stButton"] > button {
                background: #dc2626;
                border: 1px solid #b91c1c;
                color: #ffffff;
                font-weight: 600;
            }
            .st-key-track-info-reset div[data-testid="stButton"] > button:hover {
                background: #b91c1c;
                border-color: #991b1b;
                color: #ffffff;
            }
            .st-key-track-info-export-info div[data-testid="stDownloadButton"] > button:not(:disabled) {
                background: #16a34a;
                border: 1px solid #15803d;
                color: #ffffff;
                font-weight: 600;
            }
            .st-key-track-info-export-info div[data-testid="stDownloadButton"] > button:not(:disabled):hover {
                background: #15803d;
                border-color: #166534;
                color: #ffffff;
            }
            .instrument-autofill-label {
                color: var(--text-color, inherit);
                font-size: 0.92rem;
                line-height: 1.3;
                opacity: 0.82;
                padding-top: 0.2rem;
                white-space: nowrap;
            }
            .instrument-autofill-prefix {
                color: var(--text-color, inherit);
                opacity: 0.82;
                white-space: nowrap;
            }
            .meter-slash {
                color: #d1d5db;
                font-size: 1rem;
                line-height: 2.4rem;
                text-align: center;
                padding-top: 0.2rem;
            }
        """
    st.set_page_config(page_title=APP_TITLE, layout="wide")
    st.markdown(
        base_css + body_font_css + title_font_css + track_header_css + "\n</style>",
        unsafe_allow_html=True,
    )


def compact_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def unique_in_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    clean_values: list[str] = []

    for raw_value in values:
        value = compact_text(raw_value)
        value_key = value.casefold()
        if not value or value_key in {"instrument", "instruments", "instrumentation"}:
            continue
        if value_key in seen:
            continue
        seen.add(value_key)
        clean_values.append(value)

    return sorted(clean_values, key=str.casefold)


def clean_values_in_source_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    clean_values: list[str] = []

    for raw_value in values:
        value = compact_text(raw_value)
        value_key = value.casefold()
        if not value or value_key in {"instrument", "instruments", "instrumentation"}:
            continue
        if value_key in seen:
            continue
        seen.add(value_key)
        clean_values.append(value)

    return clean_values


def read_single_column_workbook(workbook_path: Path) -> list[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    values: list[str] = []

    for row in sheet.iter_rows(values_only=True):
        for cell_value in row:
            if cell_value is not None:
                values.append(str(cell_value))

    workbook.close()
    return unique_in_order(values)


def read_single_column_workbook_in_source_order(workbook_path: Path) -> list[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    values: list[str] = []

    for row in sheet.iter_rows(values_only=True):
        for cell_value in row:
            if cell_value is not None:
                values.append(str(cell_value))

    workbook.close()
    return clean_values_in_source_order(values)


@st.cache_data
def load_instrument_options(taxonomy_modified_at: int) -> list[str]:
    _ = taxonomy_modified_at
    return read_single_column_workbook(INSTRUMENT_TAXONOMY_PATH)


@st.cache_data
def load_vocal_options(vocal_sublist_modified_at: int) -> list[str]:
    _ = vocal_sublist_modified_at
    return read_single_column_workbook(VOCAL_SUBLIST_PATH)


@st.cache_data
def load_key_options(key_selection_modified_at: int) -> list[str]:
    _ = key_selection_modified_at
    key_options = read_single_column_workbook_in_source_order(KEY_SELECTION_PATH)
    key_options = [option for option in key_options if option != "No Key"]
    return ["No Key", *key_options]


@st.cache_data
def load_binary_asset_data_uri(asset_path: str, asset_modified_at: int) -> str:
    _ = asset_modified_at
    path = Path(asset_path)
    mime_type = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
        ".ttf": "font/ttf",
        ".otf": "font/otf",
        ".woff": "font/woff",
        ".woff2": "font/woff2",
    }.get(path.suffix.lower(), "application/octet-stream")
    encoded_asset = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded_asset}"


def render_logo_header() -> None:
    if not LOGO_PATH.exists():
        return

    logo_data_uri = load_binary_asset_data_uri(
        str(LOGO_PATH),
        LOGO_PATH.stat().st_mtime_ns,
    )
    st.markdown(
        (
            '<div class="app-logo-wrap">'
            f'<img src="{logo_data_uri}" alt="Warner Chappell Production Music" '
            'class="app-logo">'
            "</div>"
        ),
        unsafe_allow_html=True,
    )


def render_page_help() -> None:
    tooltip_items = "".join(f"<li>{item}</li>" for item in HELP_TOOLTIP_ITEMS)
    st.markdown(
        (
            '<div class="page-help-row">'
            '<div class="page-help">'
            '<div class="page-help-icon" aria-label="Track Info help">?</div>'
            f'<div class="page-help-tooltip"><ul>{tooltip_items}</ul></div>'
            "</div>"
            "</div>"
        ),
        unsafe_allow_html=True,
    )


def includes_vocals(selected_instruments: list[str]) -> bool:
    return any(instrument.casefold() == "vocals" for instrument in selected_instruments)


def ensure_track_state(
    track_number: int,
    instrument_options: list[str],
    vocal_options: list[str],
) -> None:
    composer_count_key = f"track_{track_number}_composer_count"
    st.session_state.setdefault(composer_count_key, 1)

    instrumentation_key = f"track_{track_number}_instrumentation"
    selected_instruments = st.session_state.get(instrumentation_key, [])
    if selected_instruments:
        st.session_state[instrumentation_key] = [
            instrument
            for instrument in selected_instruments
            if instrument in instrument_options
        ]
        selected_instruments = st.session_state[instrumentation_key]

    vocal_sublist_key = f"track_{track_number}_vocal_sublist"
    selected_vocal_options = st.session_state.get(vocal_sublist_key, [])
    if selected_vocal_options:
        st.session_state[vocal_sublist_key] = [
            vocal_option
            for vocal_option in selected_vocal_options
            if vocal_option in vocal_options
        ]

    featured_key = f"track_{track_number}_featured_instrument"
    featured_value = st.session_state.get(featured_key, "")
    if featured_value and featured_value not in st.session_state.get(instrumentation_key, []):
        st.session_state[featured_key] = ""


def increment_composer_count(track_number: int) -> None:
    key = f"track_{track_number}_composer_count"
    st.session_state[key] = int(st.session_state.get(key, 1)) + 1


def decrement_composer_count(track_number: int) -> None:
    key = f"track_{track_number}_composer_count"
    st.session_state[key] = max(1, int(st.session_state.get(key, 1)) - 1)


def composer_prefix(track_number: int, composer_number: int) -> str:
    return f"track_{track_number}_composer_{composer_number}"


def used_instrument_suggestions(
    track_count: int,
    live_track_number: int | None = None,
    live_selected_instruments: list[str] | None = None,
) -> list[str]:
    suggestions: list[str] = []
    seen: set[str] = set()

    for track_number in range(1, track_count + 1):
        if track_number == live_track_number and live_selected_instruments is not None:
            track_instruments = live_selected_instruments
        else:
            track_instruments = st.session_state.get(
                f"track_{track_number}_instrumentation",
                [],
            )

        for instrument in track_instruments:
            instrument_key = instrument.casefold()
            if instrument_key in seen:
                continue
            seen.add(instrument_key)
            suggestions.append(instrument)

    return suggestions


def merge_instruments(existing: list[str], additions: list[str]) -> list[str]:
    merged = list(existing)
    seen = {instrument.casefold() for instrument in merged}

    for instrument in additions:
        instrument_key = instrument.casefold()
        if instrument_key in seen:
            continue
        seen.add(instrument_key)
        merged.append(instrument)

    return merged


def add_instrument_suggestion(track_number: int, instrument_name: str) -> None:
    instrumentation_key = f"track_{track_number}_instrumentation"
    current_instruments = st.session_state.get(instrumentation_key, [])
    st.session_state[instrumentation_key] = merge_instruments(
        current_instruments,
        [instrument_name],
    )


def add_all_instrument_suggestions(track_number: int, suggestions: list[str]) -> None:
    instrumentation_key = f"track_{track_number}_instrumentation"
    current_instruments = st.session_state.get(instrumentation_key, [])
    st.session_state[instrumentation_key] = merge_instruments(
        current_instruments,
        suggestions,
    )


def render_instrument_autofill(track_number: int, suggestions: list[str]) -> None:
    suggestion_items = ["ALL", *suggestions] if suggestions else []
    if not suggestion_items:
        st.markdown(
            '<div class="instrument-autofill-label">Autofill Instruments:</div>',
            unsafe_allow_html=True,
        )
        return

    suggestion_specs: list[tuple[str, str, int]] = []
    for index, item in enumerate(suggestion_items):
        display_label = f"{item}," if index < len(suggestion_items) - 1 else item
        char_weight = len(display_label) + 2
        suggestion_specs.append((item, display_label, char_weight))

    first_row_budget = 160
    later_row_budget = 160
    rows: list[list[tuple[str, str, int]]] = []
    current_row: list[tuple[str, str, int]] = []
    current_budget = first_row_budget
    current_width = 0

    for spec in suggestion_specs:
        _, _, width = spec
        if current_row and current_width + width > current_budget:
            rows.append(current_row)
            current_row = []
            current_width = 0.0
            current_budget = later_row_budget
        current_row.append(spec)
        current_width += width

    if current_row:
        rows.append(current_row)

    for row_index, row_items in enumerate(rows):
        with st.container(
            horizontal=True,
            horizontal_alignment="left",
            gap="small",
            key=f"track_{track_number}_instrument_autofill_row_{row_index}",
        ):
            if row_index == 0:
                st.markdown(
                    '<div class="instrument-autofill-label">Autofill Instruments:</div>',
                    unsafe_allow_html=True,
                )

            for item, display_label, _ in row_items:
                if item == "ALL":
                    st.button(
                        display_label,
                        key=f"track_{track_number}_autofill_all_{row_index}",
                        type="tertiary",
                        use_container_width=False,
                        on_click=add_all_instrument_suggestions,
                        args=(track_number, suggestions),
                    )
                else:
                    st.button(
                        display_label,
                        key=f"track_{track_number}_autofill_{item}_{row_index}",
                        type="tertiary",
                        use_container_width=False,
                        on_click=add_instrument_suggestion,
                        args=(track_number, item),
                    )


def parse_split_value(raw_value: object) -> float | None:
    if raw_value in ("", None):
        return None

    try:
        value = float(str(raw_value).strip().removesuffix("%"))
    except (TypeError, ValueError):
        return None

    if value < 0 or value > 100:
        return None

    return value


def read_split_value(raw_value: object, default: float = 0.0) -> float:
    parsed_value = parse_split_value(raw_value)
    if parsed_value is None:
        return default

    return parsed_value


def parse_meter_components(raw_value: object) -> tuple[str, str]:
    meter_text = compact_text(raw_value)
    if not meter_text:
        return "", ""

    numerator, separator, denominator = meter_text.partition("/")
    if not separator:
        return meter_text, ""

    return compact_text(numerator), compact_text(denominator)


def format_meter_value(numerator: object, denominator: object) -> str:
    numerator_text = compact_text(numerator)
    denominator_text = compact_text(denominator)

    if not numerator_text or not denominator_text:
        return ""
    if not numerator_text.isdigit() or not denominator_text.isdigit():
        return ""

    if int(numerator_text) <= 0 or int(denominator_text) <= 0:
        return ""

    return f"{int(numerator_text)}/{int(denominator_text)}"


def clean_multiline_text(value: object) -> str:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    return "\n".join(line.rstrip() for line in text.split("\n"))


def combine_dsp_links(spotify_link: object, apple_music_link: object) -> str:
    spotify_text = compact_text(spotify_link)
    apple_text = compact_text(apple_music_link)

    if spotify_text and apple_text:
        return f"{spotify_text}, {apple_text}"
    return spotify_text or apple_text


def parse_dsp_links(raw_value: object) -> tuple[str, str]:
    value = compact_text(raw_value)
    if not value:
        return "", ""

    first_value, separator, second_value = value.partition(",")
    if separator:
        return compact_text(first_value), compact_text(second_value)

    lowered_value = value.casefold()
    if "apple" in lowered_value or "music.apple" in lowered_value:
        return "", value

    return value, ""


def read_cae_ipi_value(raw_value: object) -> int | None:
    if raw_value in ("", None):
        return None

    raw_text = str(raw_value).strip()
    if not raw_text.isdigit():
        return None

    try:
        return int(raw_text)
    except (TypeError, ValueError):
        return None


def ensure_text_value(key: str, decimal_places: int | None = None) -> None:
    if key not in st.session_state or isinstance(st.session_state[key], str):
        return

    value = st.session_state[key]
    if value is None:
        st.session_state[key] = ""
    elif decimal_places is None:
        st.session_state[key] = str(value)
    else:
        st.session_state[key] = f"{float(value):.{decimal_places}f}"


def multi_key_count_key(track_number: int) -> str:
    return f"track_{track_number}_multi_key_count"


def multi_key_slot_key(track_number: int, slot_number: int) -> str:
    return f"track_{track_number}_multi_key_slot_{slot_number}"


def normalize_multi_key_sequence(
    raw_values: list[object],
    key_options: list[str],
) -> list[str]:
    return [
        key_value
        for key_value in (compact_text(raw_value) for raw_value in raw_values)
        if key_value in key_options
    ]


def read_multi_key_slots(track_number: int, key_options: list[str]) -> list[str]:
    slot_count = max(1, int(st.session_state.get(multi_key_count_key(track_number), 1) or 1))
    return normalize_multi_key_sequence(
        [
            st.session_state.get(multi_key_slot_key(track_number, slot_number))
            for slot_number in range(1, slot_count + 1)
        ],
        key_options,
    )


def set_multi_key_slots(
    track_number: int,
    key_options: list[str],
    values: list[object],
) -> None:
    normalized_values = normalize_multi_key_sequence(values, key_options)
    slot_count = max(1, len(normalized_values))
    st.session_state[multi_key_count_key(track_number)] = slot_count

    for slot_number in range(1, slot_count + 1):
        st.session_state[multi_key_slot_key(track_number, slot_number)] = (
            normalized_values[slot_number - 1]
            if slot_number <= len(normalized_values)
            else None
        )


def ensure_multi_key_slot_state(
    track_number: int,
    key_options: list[str],
    default_values: list[object],
) -> None:
    normalized_values = normalize_multi_key_sequence(default_values, key_options)
    count_key = multi_key_count_key(track_number)
    slot_count = max(
        1,
        int(st.session_state.get(count_key, len(normalized_values) or 1) or 1),
        len(normalized_values),
    )
    st.session_state[count_key] = slot_count

    for slot_number in range(1, slot_count + 1):
        slot_key = multi_key_slot_key(track_number, slot_number)
        current_value = st.session_state.get(slot_key)
        if current_value in key_options:
            continue
        st.session_state[slot_key] = (
            normalized_values[slot_number - 1]
            if slot_number <= len(normalized_values)
            else None
        )


def read_composer_record(track_number: int, composer_number: int) -> dict[str, object]:
    prefix = composer_prefix(track_number, composer_number)
    split_default = 100.0 if composer_number == 1 else 0.0
    spotify_link = compact_text(st.session_state.get(f"{prefix}_spotify_link", ""))
    apple_music_link = compact_text(
        st.session_state.get(f"{prefix}_apple_music_link", "")
    )

    return {
        "composer": compact_text(st.session_state.get(f"{prefix}_composer", "")),
        "pro_affiliation": compact_text(
            st.session_state.get(f"{prefix}_pro_affiliation", "")
        ),
        "cae_ipi": read_cae_ipi_value(st.session_state.get(f"{prefix}_cae_ipi")),
        "split": parse_split_value(
            st.session_state.get(f"{prefix}_split", f"{split_default:.2f}")
        ),
        "spotify_link": spotify_link,
        "apple_music_link": apple_music_link,
        "dsp_link": combine_dsp_links(spotify_link, apple_music_link),
        "artist_name": compact_text(st.session_state.get(f"{prefix}_artist_name", "")),
    }


def composer_has_text(record: dict[str, object]) -> bool:
    return (
        any(compact_text(record.get(field_name, "")) for field_name in COMPOSER_TEXT_FIELDS)
        or record.get("cae_ipi") is not None
    )


def composer_label(record: dict[str, object]) -> str:
    parts = [record["composer"]]

    if record["pro_affiliation"]:
        parts.append(record["pro_affiliation"])
    if record["cae_ipi"] is not None:
        parts.append(f"CAE/IPI: {record['cae_ipi']}")
    parts.append(f"{read_split_value(record.get('split')):g}%")
    if record["artist_name"]:
        parts.append(f"Artist: {record['artist_name']}")
    if record["spotify_link"] or record["apple_music_link"]:
        parts.append("DSP link")

    return " | ".join(parts)


def build_composer_lookup(track_count: int) -> dict[str, dict[str, object]]:
    lookup: dict[str, dict[str, object]] = {}

    for track_number in range(1, track_count + 1):
        composer_count = int(
            st.session_state.get(f"track_{track_number}_composer_count", 1)
        )
        for composer_number in range(1, composer_count + 1):
            record = read_composer_record(track_number, composer_number)
            if not record["composer"]:
                continue

            lookup_key = "|".join(
                [
                    record["composer"].casefold(),
                    record["pro_affiliation"].casefold(),
                    "" if record["cae_ipi"] is None else str(record["cae_ipi"]),
                    f"{read_split_value(record.get('split')):.4f}",
                    record["spotify_link"].casefold(),
                    record["apple_music_link"].casefold(),
                    record["artist_name"].casefold(),
                ]
            )
            lookup.setdefault(lookup_key, record)

    return lookup


def apply_saved_composer(target_prefix: str, select_key: str) -> None:
    selected_key = st.session_state.get(select_key, "")
    if not selected_key:
        return

    record = st.session_state.get("_composer_lookup", {}).get(selected_key)
    if not record:
        return

    for field_name in COMPOSER_FIELDS:
        st.session_state[f"{target_prefix}_{field_name}"] = record[field_name]


def format_saved_composer_option(option_key: str) -> str:
    if not option_key:
        return "No composer info available yet"

    record = st.session_state.get("_composer_lookup", {}).get(option_key)
    if not record:
        return option_key

    return composer_label(record)


def render_header_fields() -> int:
    st.session_state.setdefault("track_count", 1)
    header_row = st.columns([2.2, 1.6, 0.45])
    with header_row[0]:
        st.text_input("Album Name:", key="album_name")
    with header_row[1]:
        st.text_input("WCPM Producer:", key="producer")
    with header_row[2]:
        track_count = st.number_input(
            "Track Count",
            min_value=1,
            max_value=200,
            step=1,
            format="%d",
            key="track_count",
        )

    return int(track_count)


def get_instrument_options() -> list[str]:
    try:
        taxonomy_modified_at = INSTRUMENT_TAXONOMY_PATH.stat().st_mtime_ns
        instrument_options = load_instrument_options(taxonomy_modified_at)
    except Exception as exc:
        st.error(f"Instrument taxonomy could not be loaded: {exc}")
        return []

    if not instrument_options:
        st.error("Instrument taxonomy is empty.")

    return instrument_options


def get_vocal_options() -> list[str]:
    try:
        vocal_sublist_modified_at = VOCAL_SUBLIST_PATH.stat().st_mtime_ns
        vocal_options = load_vocal_options(vocal_sublist_modified_at)
    except Exception as exc:
        st.error(f"Vocal sub-list could not be loaded: {exc}")
        return []

    if not vocal_options:
        st.error("Vocal sub-list is empty.")

    return vocal_options


def get_key_options() -> list[str]:
    try:
        key_selection_modified_at = KEY_SELECTION_PATH.stat().st_mtime_ns
        key_options = load_key_options(key_selection_modified_at)
    except Exception as exc:
        st.error(f"Key selection list could not be loaded: {exc}")
        return []

    if not key_options:
        st.error("Key selection list is empty.")

    return key_options


def render_track_fields(
    track_number: int,
    instrument_options: list[str],
    vocal_options: list[str],
    key_options: list[str],
) -> None:
    ensure_track_state(track_number, instrument_options, vocal_options)
    st.session_state["_composer_lookup"] = build_composer_lookup(
        int(st.session_state.get("track_count", 1))
    )

    track_title = compact_text(st.session_state.get(f"track_{track_number}_title", ""))
    expander_label = f"Track {track_number}"
    if track_title:
        expander_label = f"{expander_label} - {track_title}"

    with st.container(key=f"track-panel-{track_number}"):
        with st.expander(expander_label, expanded=True):
            key_value_key = f"track_{track_number}_key"
            multi_key_toggle_key = f"track_{track_number}_multi_key_enabled"
            multi_key_previous_key = f"track_{track_number}_multi_key_previous"
            single_key_key = f"track_{track_number}_single_key"
            multi_keys_key = f"track_{track_number}_multi_keys"
            multi_key_count_state_key = multi_key_count_key(track_number)
            meter_value_key = f"track_{track_number}_meter"
            meter_numerator_key = f"track_{track_number}_meter_numerator"
            meter_denominator_key = f"track_{track_number}_meter_denominator"
            lyrics_toggle_key = f"track_{track_number}_has_lyrics"
            lyrics_key = f"track_{track_number}_lyrics"

            current_key_text = compact_text(st.session_state.get(key_value_key, ""))
            current_key_values = [
                key_option
                for key_option in current_key_text.split(", ")
                if key_option in key_options
            ]

            if multi_key_toggle_key not in st.session_state:
                st.session_state[multi_key_toggle_key] = len(current_key_values) > 1
            if single_key_key not in st.session_state:
                st.session_state[single_key_key] = (
                    current_key_values[0] if current_key_values else None
                )
            if multi_keys_key not in st.session_state:
                st.session_state[multi_keys_key] = current_key_values
            if multi_key_count_state_key not in st.session_state:
                st.session_state[multi_key_count_state_key] = max(1, len(current_key_values))

            if st.session_state.get(single_key_key) not in key_options:
                st.session_state[single_key_key] = None
            st.session_state[multi_keys_key] = normalize_multi_key_sequence(
                st.session_state.get(multi_keys_key, []),
                key_options,
            )
            ensure_multi_key_slot_state(
                track_number,
                key_options,
                st.session_state.get(multi_keys_key, []),
            )

            meter_numerator, meter_denominator = parse_meter_components(
                st.session_state.get(meter_value_key, "")
            )
            st.session_state.setdefault(meter_numerator_key, meter_numerator)
            st.session_state.setdefault(meter_denominator_key, meter_denominator)
            st.session_state.setdefault(lyrics_toggle_key, False)
            st.session_state.setdefault(lyrics_key, "")

            label_cols = st.columns([2.75, 0.55, 1.35, 0.9, 0.85, 1.0])
            with label_cols[0]:
                st.markdown("Title:")
            with label_cols[1]:
                st.markdown("BPM:")
            with label_cols[2]:
                st.markdown("Key:")
            with label_cols[3]:
                st.markdown("Multi-Key?")
            with label_cols[4]:
                st.markdown("Meter:")
            with label_cols[5]:
                st.markdown("Has Lyrics?")

            multi_key_enabled = st.session_state.get(multi_key_toggle_key, False)
            previous_multi_key_enabled = st.session_state.get(
                multi_key_previous_key,
                multi_key_enabled,
            )
            input_cols = st.columns([2.75, 0.55, 1.35, 0.9, 0.85, 1.0])
            with input_cols[0]:
                st.text_input(
                    "Title:",
                    key=f"track_{track_number}_title",
                    label_visibility="collapsed",
                )
            with input_cols[1]:
                st.number_input(
                    "BPM:",
                    min_value=0,
                    max_value=400,
                    step=1,
                    format="%d",
                    key=f"track_{track_number}_bpm",
                    label_visibility="collapsed",
                )
            with input_cols[2]:
                multi_key_enabled = st.session_state.get(multi_key_toggle_key, False)
                if multi_key_enabled != previous_multi_key_enabled:
                    if multi_key_enabled:
                        existing_multi_keys = normalize_multi_key_sequence(
                            st.session_state.get(multi_keys_key, []),
                            key_options,
                        )
                        if existing_multi_keys:
                            set_multi_key_slots(
                                track_number,
                                key_options,
                                existing_multi_keys,
                            )
                        else:
                            current_single_key = st.session_state.get(single_key_key)
                            set_multi_key_slots(
                                track_number,
                                key_options,
                                [current_single_key] if current_single_key in key_options else [],
                            )
                    else:
                        current_multi_keys = read_multi_key_slots(track_number, key_options)
                        st.session_state[multi_keys_key] = current_multi_keys
                        st.session_state[single_key_key] = (
                            current_multi_keys[0] if current_multi_keys else None
                        )
                    st.session_state[multi_key_previous_key] = multi_key_enabled

                if multi_key_enabled:
                    slot_count = max(
                        1,
                        int(st.session_state.get(multi_key_count_state_key, 1) or 1),
                    )
                    for slot_number in range(1, slot_count + 1):
                        st.selectbox(
                            f"Key {slot_number}:",
                            options=key_options,
                            index=None,
                            placeholder="Select key",
                            key=multi_key_slot_key(track_number, slot_number),
                            label_visibility="collapsed",
                        )

                    selected_keys = read_multi_key_slots(track_number, key_options)
                    st.session_state[multi_keys_key] = selected_keys
                    st.session_state[key_value_key] = ", ".join(selected_keys)

                    key_control_cols = st.columns([1, 1])
                    with key_control_cols[0]:
                        if st.button(
                            "+",
                            key=f"track_{track_number}_add_multi_key_slot",
                            help="Add another key to the multi-key sequence",
                            use_container_width=True,
                        ):
                            st.session_state[multi_key_count_state_key] = slot_count + 1
                            st.session_state[multi_key_slot_key(track_number, slot_count + 1)] = None
                            st.rerun()
                    with key_control_cols[1]:
                        if st.button(
                            "-",
                            key=f"track_{track_number}_remove_multi_key_slot",
                            help="Remove the last key from the multi-key sequence",
                            disabled=slot_count <= 1,
                            use_container_width=True,
                        ):
                            st.session_state.pop(
                                multi_key_slot_key(track_number, slot_count),
                                None,
                            )
                            st.session_state[multi_key_count_state_key] = max(1, slot_count - 1)
                            st.rerun()
                else:
                    st.selectbox(
                        "Key:",
                        options=key_options,
                        index=None,
                        placeholder="Select key",
                        key=single_key_key,
                        label_visibility="collapsed",
                    )
                    st.session_state[key_value_key] = (
                        st.session_state.get(single_key_key) or ""
                    )
            with input_cols[3]:
                st.checkbox(
                    "Multi-Key?",
                    key=multi_key_toggle_key,
                    label_visibility="collapsed",
                )
            with input_cols[4]:
                meter_cols = st.columns([1, 0.12, 1])
                with meter_cols[0]:
                    st.text_input(
                        "Meter numerator:",
                        key=meter_numerator_key,
                        label_visibility="collapsed",
                        placeholder="-",
                        max_chars=3,
                    )
                with meter_cols[1]:
                    st.markdown('<div class="meter-slash">/</div>', unsafe_allow_html=True)
                with meter_cols[2]:
                    st.text_input(
                        "Meter denominator:",
                        key=meter_denominator_key,
                        label_visibility="collapsed",
                        placeholder="-",
                        max_chars=3,
                    )

                st.session_state[meter_value_key] = format_meter_value(
                    st.session_state.get(meter_numerator_key, ""),
                    st.session_state.get(meter_denominator_key, ""),
                )
            with input_cols[5]:
                st.checkbox(
                    "Has Lyrics?",
                    key=lyrics_toggle_key,
                    label_visibility="collapsed",
                )

            st.session_state[multi_key_previous_key] = st.session_state.get(
                multi_key_toggle_key,
                False,
            )

            instrument_col, featured_col = st.columns([4.0, 0.9])
            instrumentation_key = f"track_{track_number}_instrumentation"
            featured_key = f"track_{track_number}_featured_instrument"

            with instrument_col:
                selected_instruments = st.multiselect(
                    "Instrumentation:",
                    options=instrument_options,
                    key=instrumentation_key,
                )

            with featured_col:
                if selected_instruments:
                    if st.session_state.get(featured_key) not in selected_instruments:
                        st.session_state[featured_key] = None

                    st.selectbox(
                        "Featured Instrument:",
                        options=selected_instruments,
                        index=None,
                        placeholder="Select featured instrument",
                        key=featured_key,
                    )
                else:
                    st.session_state[featured_key] = None
                    st.selectbox(
                        "Featured Instrument:",
                        options=[],
                        placeholder="Select instrumentation first",
                        disabled=True,
                        key=featured_key,
                    )

            render_instrument_autofill(
                track_number,
                used_instrument_suggestions(
                    int(st.session_state.get("track_count", 1)),
                    live_track_number=track_number,
                    live_selected_instruments=selected_instruments,
                ),
            )

            if includes_vocals(selected_instruments):
                st.multiselect(
                    "Vocal Sub-list:",
                    options=vocal_options,
                    key=f"track_{track_number}_vocal_sublist",
                )

            if st.session_state.get(lyrics_toggle_key, False):
                st.text_area(
                    "Lyrics:",
                    key=lyrics_key,
                    height=180,
                    placeholder="Enter lyrics",
                )

            st.divider()
            composer_count = int(st.session_state.get(f"track_{track_number}_composer_count", 1))
            composer_lookup = st.session_state.get("_composer_lookup", {})
            saved_options = [""] + list(composer_lookup.keys())

            for composer_number in range(1, composer_count + 1):
                prefix = composer_prefix(track_number, composer_number)
                saved_key = f"{prefix}_saved"
                if st.session_state.get(saved_key, "") not in saved_options:
                    st.session_state[saved_key] = ""

                st.selectbox(
                    f"Autofill Composer Info {composer_number}:",
                    options=saved_options,
                    format_func=format_saved_composer_option,
                    key=saved_key,
                    on_change=apply_saved_composer,
                    args=(prefix, saved_key),
                    disabled=not composer_lookup,
                )

                split_key = f"{prefix}_split"
                st.session_state.setdefault(
                    split_key,
                    f"{100.0 if composer_number == 1 else 0.0:.2f}",
                )
                ensure_text_value(f"{prefix}_cae_ipi")
                ensure_text_value(split_key, decimal_places=2)

                composer_col, pro_col, cae_col, split_col, spotify_col, apple_col, artist_col = st.columns(
                    [2, 0.9, 0.8, 0.9, 1, 1, 1.8]
                )
                with composer_col:
                    st.text_input("Composer:", key=f"{prefix}_composer")
                with pro_col:
                    st.text_input("PRO Affiliation:", key=f"{prefix}_pro_affiliation")
                with cae_col:
                    st.text_input(
                        "CAE/IPI",
                        key=f"{prefix}_cae_ipi",
                    )
                with split_col:
                    st.text_input(
                        "Splits:",
                        key=split_key,
                    )
                with spotify_col:
                    st.text_input(
                        "DSP Links - Spotify:",
                        key=f"{prefix}_spotify_link",
                    )
                with apple_col:
                    st.text_input(
                        "Apple Music:",
                        key=f"{prefix}_apple_music_link",
                    )
                with artist_col:
                    st.text_input(
                        "Artist Name (Optional):",
                        key=f"{prefix}_artist_name",
                    )

            add_col, remove_col, _ = st.columns([1, 1, 4])
            with add_col:
                st.button(
                    "Add Composer",
                    key=f"track_{track_number}_add_composer",
                    on_click=increment_composer_count,
                    args=(track_number,),
                    use_container_width=True,
                )
            with remove_col:
                st.button(
                    "Remove Composer",
                    key=f"track_{track_number}_remove_composer",
                    on_click=decrement_composer_count,
                    args=(track_number,),
                    disabled=composer_count <= 1,
                    use_container_width=True,
                )


def collect_tracks(track_count: int) -> list[dict[str, object]]:
    tracks: list[dict[str, object]] = []

    for track_number in range(1, track_count + 1):
        composer_count = int(
            st.session_state.get(f"track_{track_number}_composer_count", 1)
        )
        selected_instruments = st.session_state.get(f"track_{track_number}_instrumentation", [])
        composers = [
            read_composer_record(track_number, composer_number)
            for composer_number in range(1, composer_count + 1)
        ]
        meter_numerator = compact_text(
            st.session_state.get(f"track_{track_number}_meter_numerator", "")
        )
        meter_denominator = compact_text(
            st.session_state.get(f"track_{track_number}_meter_denominator", "")
        )

        bpm_value = st.session_state.get(f"track_{track_number}_bpm", 0)
        tracks.append(
            {
                "album_name": compact_text(st.session_state.get("album_name", "")),
                "producer": compact_text(st.session_state.get("producer", "")),
                "track_number": track_number,
                "track_title": compact_text(
                    st.session_state.get(f"track_{track_number}_title", "")
                ),
                "bpm": int(bpm_value) if bpm_value else "",
                "key": compact_text(st.session_state.get(f"track_{track_number}_key", "")),
                "meter": format_meter_value(meter_numerator, meter_denominator),
                "meter_numerator": meter_numerator,
                "meter_denominator": meter_denominator,
                "instrumentation": ", ".join(
                    selected_instruments
                ),
                "has_lyrics": bool(st.session_state.get(f"track_{track_number}_has_lyrics", False)),
                "lyrics": clean_multiline_text(
                    st.session_state.get(f"track_{track_number}_lyrics", "")
                ),
                "vocal_sublist": ", ".join(
                    st.session_state.get(f"track_{track_number}_vocal_sublist", [])
                    if includes_vocals(selected_instruments)
                    else []
                ),
                "featured_instrument": compact_text(
                    st.session_state.get(
                        f"track_{track_number}_featured_instrument", ""
                    )
                ),
                "composers": composers,
            }
        )

    return tracks


def composer_headers(max_composers: int) -> list[str]:
    headers: list[str] = []
    for composer_number in range(1, max_composers + 1):
        headers.extend(
            [
                f"Composer {composer_number}",
                f"PRO Affiliation {composer_number}",
                f"CAE/IPI {composer_number}",
                f"Splits {composer_number} (%)",
                f"DSP Links (Spotify, Apple Music) {composer_number}",
                f"Artist Name {composer_number} (Optional)",
            ]
        )
    return headers


def track_to_row(track: dict[str, object], max_composers: int) -> list[object]:
    row = [
        track["album_name"],
        track["producer"],
        track["track_number"],
        track["track_title"],
        track["bpm"],
        track["key"],
        track["meter"],
        track["instrumentation"],
        track["vocal_sublist"],
        track["featured_instrument"],
    ]

    composers = track["composers"]
    for composer_index in range(max_composers):
        if composer_index < len(composers):
            record = composers[composer_index]
            row.extend(
                [
                    record["composer"],
                    record["pro_affiliation"],
                    record["cae_ipi"] if record["cae_ipi"] is not None else "",
                    read_split_value(record["split"]) / 100,
                    combine_dsp_links(
                        record.get("spotify_link", ""),
                        record.get("apple_music_link", ""),
                    ),
                    record["artist_name"],
                ]
            )
        else:
            row.extend(["", "", "", "", "", ""])

    return row


def build_excel_workbook(tracks: list[dict[str, object]]) -> bytes:
    max_composers = max(1, *(len(track["composers"]) for track in tracks))
    headers = BASE_HEADERS + composer_headers(max_composers)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Track Info"

    sheet.append(headers)
    for track in tracks:
        sheet.append(track_to_row(track, max_composers))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = BODY_ALIGNMENT

    column_widths = {
        "A": 22,
        "B": 20,
        "C": 12,
        "D": 34,
        "E": 10,
        "F": 12,
        "G": 12,
        "H": 42,
        "I": 28,
        "J": 24,
    }
    for column_letter, width in column_widths.items():
        sheet.column_dimensions[column_letter].width = width

    for col_idx in range(len(BASE_HEADERS) + 1, len(headers) + 1):
        header = headers[col_idx - 1]
        if header.startswith("DSP Links (Spotify, Apple Music) "):
            width = 34
        elif header.startswith("CAE/IPI "):
            width = 16
        elif header.startswith("Splits "):
            width = 14
        elif "Artist Name" in header:
            width = 24
        else:
            width = 22
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    for col_idx, header in enumerate(headers, start=1):
        if header.startswith("Splits "):
            for row_idx in range(2, sheet.max_row + 1):
                sheet.cell(row=row_idx, column=col_idx).number_format = "0.00%"
        elif header.startswith("CAE/IPI "):
            for row_idx in range(2, sheet.max_row + 1):
                sheet.cell(row=row_idx, column=col_idx).number_format = "0"

    sheet.row_dimensions[1].height = 34

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def tracks_with_lyrics(tracks: list[dict[str, object]]) -> list[dict[str, object]]:
    return [
        track
        for track in tracks
        if track.get("has_lyrics") and track.get("lyrics")
    ]


def docx_text_run_xml(text: str, *, bold: bool = False, font_size_half_points: int = 22) -> str:
    run_properties = [f'<w:sz w:val="{font_size_half_points}"/>']
    if bold:
        run_properties.insert(0, "<w:b/>")
    escaped_text = escape(text)
    return (
        "<w:r>"
        f"<w:rPr>{''.join(run_properties)}</w:rPr>"
        f'<w:t xml:space="preserve">{escaped_text}</w:t>'
        "</w:r>"
    )


def lyrics_to_paragraph_xml(lyrics_text: str) -> str:
    lyric_lines = lyrics_text.split("\n")
    paragraphs: list[str] = []

    for lyric_line in lyric_lines:
        if lyric_line:
            paragraphs.append(
                f"<w:p>{docx_text_run_xml(lyric_line, font_size_half_points=22)}</w:p>"
            )
        else:
            paragraphs.append("<w:p/>")

    return "".join(paragraphs)


def build_lyrics_docx(tracks: list[dict[str, object]]) -> bytes:
    lyric_tracks = tracks_with_lyrics(tracks)
    body_parts: list[str] = []

    for index, track in enumerate(lyric_tracks):
        title_text = f"Title: {track['track_title']}"
        body_parts.append(
            "<w:p>"
            f"{docx_text_run_xml(title_text, bold=True, font_size_half_points=28)}"
            "</w:p>"
        )
        body_parts.append(lyrics_to_paragraph_xml(track["lyrics"]))
        if index < len(lyric_tracks) - 1:
            body_parts.append("<w:p/>")

    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"
 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
 xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
 xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:w10="urn:schemas-microsoft-com:office:word"
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"
 xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
 xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk"
 xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
 mc:Ignorable="w14 wp14">
  <w:body>
    {''.join(body_parts)}
    <w:sectPr>
      <w:pgSz w:w="12240" w:h="15840"/>
      <w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440" w:header="720" w:footer="720" w:gutter="0"/>
    </w:sectPr>
  </w:body>
</w:document>
"""

    content_types_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>
"""

    rels_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""

    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as docx_zip:
        docx_zip.writestr("[Content_Types].xml", content_types_xml)
        docx_zip.writestr("_rels/.rels", rels_xml)
        docx_zip.writestr("word/document.xml", document_xml)
    output.seek(0)
    return output.getvalue()


def paragraph_text_from_docx_xml(paragraph: ET.Element) -> str:
    parts: list[str] = []
    for node in paragraph.iter():
        if node.tag == f"{{{WORDPROCESSINGML_NS}}}t":
            parts.append(node.text or "")
        elif node.tag == f"{{{WORDPROCESSINGML_NS}}}br":
            parts.append("\n")
        elif node.tag == f"{{{WORDPROCESSINGML_NS}}}tab":
            parts.append("\t")
    return "".join(parts)


def finalize_lyrics_lines(lines: list[str]) -> str:
    while lines and not lines[-1].strip():
        lines.pop()
    return clean_multiline_text("\n".join(lines))


def parse_imported_lyrics_docx(file_bytes: bytes) -> list[dict[str, str]]:
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as docx_zip:
            document_xml = docx_zip.read("word/document.xml")
    except KeyError as exc:
        raise ValueError("This lyrics file is missing its main document content.") from exc
    except zipfile.BadZipFile as exc:
        raise ValueError("This file is not a valid Word document.") from exc

    try:
        root = ET.fromstring(document_xml)
    except ET.ParseError as exc:
        raise ValueError("This lyrics file could not be parsed.") from exc

    body = root.find(f".//{{{WORDPROCESSINGML_NS}}}body")
    if body is None:
        raise ValueError("This lyrics file does not contain any body content.")

    lyrics_entries: list[dict[str, str]] = []
    current_title = ""
    current_lines: list[str] = []

    for paragraph in body.findall(f"{{{WORDPROCESSINGML_NS}}}p"):
        paragraph_text = paragraph_text_from_docx_xml(paragraph)
        compact_paragraph_text = compact_text(paragraph_text)

        if compact_paragraph_text.startswith("Title: "):
            if current_title:
                lyrics_entries.append(
                    {
                        "track_title": current_title,
                        "lyrics": finalize_lyrics_lines(current_lines),
                    }
                )
            current_title = compact_text(compact_paragraph_text.removeprefix("Title: "))
            current_lines = []
        elif current_title:
            current_lines.append(paragraph_text.rstrip())

    if current_title:
        lyrics_entries.append(
            {
                "track_title": current_title,
                "lyrics": finalize_lyrics_lines(current_lines),
            }
        )

    if not lyrics_entries:
        raise ValueError("No lyrics tracks were found in that Word document.")

    return lyrics_entries


def merge_imported_lyrics(
    imported_workbook: dict[str, object],
    imported_lyrics: list[dict[str, str]],
) -> tuple[int, list[str]]:
    lyrics_by_title: dict[str, list[str]] = {}
    for lyrics_entry in imported_lyrics:
        title_key = compact_text(lyrics_entry["track_title"]).casefold()
        if not title_key:
            continue
        lyrics_by_title.setdefault(title_key, []).append(lyrics_entry["lyrics"])

    matched_count = 0
    for track in imported_workbook["tracks"]:
        title_key = compact_text(track["track_title"]).casefold()
        lyrics_matches = lyrics_by_title.get(title_key, [])
        if not lyrics_matches:
            continue

        track["lyrics"] = lyrics_matches.pop(0)
        track["has_lyrics"] = bool(track["lyrics"])
        matched_count += 1

    unmatched_titles: list[str] = []
    for lyrics_entry in imported_lyrics:
        title_key = compact_text(lyrics_entry["track_title"]).casefold()
        if title_key in lyrics_by_title and lyrics_by_title[title_key]:
            unmatched_titles.append(lyrics_entry["track_title"])
            lyrics_by_title[title_key].pop(0)

    return matched_count, unmatched_titles


def normalize_import_header(raw_header: object) -> str:
    header = compact_text(raw_header)
    if not header:
        return ""
    if header == "Producer":
        return "WCPM Producer"

    dsp_match = re.fullmatch(r"DSP Link (\d+)", header)
    if dsp_match:
        return f"DSP Links (Spotify, Apple Music) {dsp_match.group(1)}"

    return header


def split_exported_cell(raw_value: object) -> list[str]:
    value = compact_text(raw_value)
    if not value:
        return []
    return [part for part in (compact_text(piece) for piece in value.split(",")) if part]


def read_integer_cell(raw_value: object) -> int | None:
    if raw_value in ("", None):
        return None

    try:
        return int(float(str(raw_value).strip()))
    except (TypeError, ValueError):
        return None


def read_imported_split_value(raw_value: object) -> float | None:
    if raw_value in ("", None):
        return None

    if isinstance(raw_value, (int, float)):
        numeric_value = float(raw_value)
    else:
        try:
            numeric_value = float(str(raw_value).strip().removesuffix("%"))
        except (TypeError, ValueError):
            return None

    if 0 <= numeric_value <= 1.00001:
        numeric_value *= 100.0

    if numeric_value < 0 or numeric_value > 100:
        return None

    return round(numeric_value, 2)


def composer_numbers_from_headers(headers: list[str]) -> list[int]:
    composer_numbers: set[int] = set()
    header_patterns = (
        r"Composer (\d+)",
        r"PRO Affiliation (\d+)",
        r"CAE/IPI (\d+)",
        r"Splits (\d+) \(%\)",
        r"DSP Links \(Spotify, Apple Music\) (\d+)",
        r"Artist Name (\d+) \(Optional\)",
    )

    for header in headers:
        for pattern in header_patterns:
            match = re.fullmatch(pattern, header)
            if match:
                composer_numbers.add(int(match.group(1)))
                break

    return sorted(composer_numbers)


def imported_row_value(
    row_values: tuple[object, ...],
    header_index: dict[str, int],
    header_name: str,
) -> object:
    column_index = header_index.get(header_name)
    if column_index is None or column_index >= len(row_values):
        return None
    return row_values[column_index]


def imported_composer_record(
    row_values: tuple[object, ...],
    header_index: dict[str, int],
    composer_number: int,
) -> dict[str, object]:
    imported_dsp_link = imported_row_value(
        row_values,
        header_index,
        f"DSP Links (Spotify, Apple Music) {composer_number}",
    )
    spotify_link, apple_music_link = parse_dsp_links(imported_dsp_link)

    return {
        "composer": compact_text(
            imported_row_value(row_values, header_index, f"Composer {composer_number}")
        ),
        "pro_affiliation": compact_text(
            imported_row_value(
                row_values,
                header_index,
                f"PRO Affiliation {composer_number}",
            )
        ),
        "cae_ipi": read_cae_ipi_value(
            imported_row_value(row_values, header_index, f"CAE/IPI {composer_number}")
        ),
        "split": read_imported_split_value(
            imported_row_value(
                row_values,
                header_index,
                f"Splits {composer_number} (%)",
            )
        ),
        "spotify_link": spotify_link,
        "apple_music_link": apple_music_link,
        "dsp_link": combine_dsp_links(spotify_link, apple_music_link),
        "artist_name": compact_text(
            imported_row_value(
                row_values,
                header_index,
                f"Artist Name {composer_number} (Optional)",
            )
        ),
    }


def track_has_import_data(track: dict[str, object]) -> bool:
    return any(
        [
            track["album_name"],
            track["producer"],
            track["track_title"],
            track["bpm"] is not None,
            track["key"],
            track["meter"],
            bool(track["instrumentation"]),
            bool(track["vocal_sublist"]),
            track["featured_instrument"],
            bool(track["composers"]),
        ]
    )


def parse_imported_workbook(file_bytes: bytes) -> dict[str, object]:
    workbook = load_workbook(BytesIO(file_bytes), data_only=True)

    try:
        sheet = workbook["Track Info"] if "Track Info" in workbook.sheetnames else workbook.active
        header_row = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        headers = [normalize_import_header(cell_value) for cell_value in header_row]
        header_index = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        if not any(header in header_index for header in BASE_HEADERS):
            raise ValueError("This file does not look like a Track Info export.")

        composer_numbers = composer_numbers_from_headers(headers)
        tracks: list[dict[str, object]] = []

        for row_values in sheet.iter_rows(min_row=2, values_only=True):
            track = {
                "album_name": compact_text(
                    imported_row_value(row_values, header_index, "Album Name")
                ),
                "producer": compact_text(
                    imported_row_value(row_values, header_index, "WCPM Producer")
                ),
                "track_number": read_integer_cell(
                    imported_row_value(row_values, header_index, "Track Number")
                ),
                "track_title": compact_text(
                    imported_row_value(row_values, header_index, "Track Title")
                ),
                "bpm": read_integer_cell(
                    imported_row_value(row_values, header_index, "BPM")
                ),
                "key": compact_text(imported_row_value(row_values, header_index, "Key")),
                "meter": compact_text(
                    imported_row_value(row_values, header_index, "Meter")
                ),
                "instrumentation": split_exported_cell(
                    imported_row_value(row_values, header_index, "Instrumentation")
                ),
                "vocal_sublist": split_exported_cell(
                    imported_row_value(row_values, header_index, "Vocal Sub-list")
                ),
                "featured_instrument": compact_text(
                    imported_row_value(row_values, header_index, "Featured Instrument")
                ),
                "has_lyrics": False,
                "lyrics": "",
                "composers": [],
            }

            for composer_number in composer_numbers:
                record = imported_composer_record(
                    row_values,
                    header_index,
                    composer_number,
                )
                if composer_has_text(record) or record["split"] is not None:
                    track["composers"].append(record)

            if track_has_import_data(track):
                tracks.append(track)

        if not tracks:
            raise ValueError("No tracks were found in that workbook.")
        if len(tracks) > 200:
            raise ValueError("This workbook has more than 200 tracks, which exceeds the app limit.")

        tracks.sort(
            key=lambda track: (
                track["track_number"] is None,
                track["track_number"] or 0,
            )
        )

        album_name = next((track["album_name"] for track in tracks if track["album_name"]), "")
        producer = next((track["producer"] for track in tracks if track["producer"]), "")

        return {
            "album_name": album_name,
            "producer": producer,
            "tracks": tracks,
        }
    finally:
        workbook.close()


def build_import_state(imported_workbook: dict[str, object]) -> dict[str, object]:
    tracks: list[dict[str, object]] = list(imported_workbook["tracks"])
    state_updates: dict[str, object] = {
        "album_name": imported_workbook["album_name"],
        "producer": imported_workbook["producer"],
        "track_count": len(tracks),
    }

    for track_number, track in enumerate(tracks, start=1):
        key_text = compact_text(track["key"])
        key_values = [key for key in key_text.split(", ") if key]
        multi_key_enabled = len(key_values) > 1
        meter_numerator, meter_denominator = parse_meter_components(track["meter"])
        state_updates.update(
            {
                f"track_{track_number}_title": track["track_title"],
                f"track_{track_number}_bpm": track["bpm"] or 0,
                f"track_{track_number}_key": key_text,
                f"track_{track_number}_meter": track["meter"],
                f"track_{track_number}_meter_numerator": meter_numerator,
                f"track_{track_number}_meter_denominator": meter_denominator,
                f"track_{track_number}_instrumentation": list(track["instrumentation"]),
                f"track_{track_number}_vocal_sublist": list(track["vocal_sublist"]),
                f"track_{track_number}_featured_instrument": (
                    track["featured_instrument"] or None
                ),
                f"track_{track_number}_has_lyrics": bool(track.get("has_lyrics")),
                f"track_{track_number}_lyrics": track.get("lyrics", ""),
                f"track_{track_number}_multi_key_enabled": multi_key_enabled,
                f"track_{track_number}_multi_key_previous": multi_key_enabled,
                f"track_{track_number}_multi_key_count": max(1, len(key_values)),
                f"track_{track_number}_single_key": (
                    key_values[0] if len(key_values) == 1 else None
                ),
                f"track_{track_number}_multi_keys": key_values,
            }
        )

        composers = list(track["composers"]) or [
            {
                "composer": "",
                "pro_affiliation": "",
                "cae_ipi": None,
                "split": 100.0,
                "spotify_link": "",
                "apple_music_link": "",
                "dsp_link": "",
                "artist_name": "",
            }
        ]
        state_updates[f"track_{track_number}_composer_count"] = len(composers)

        for composer_number, record in enumerate(composers, start=1):
            prefix = composer_prefix(track_number, composer_number)
            split_default = 100.0 if composer_number == 1 else 0.0
            state_updates.update(
                {
                    f"{prefix}_saved": "",
                    f"{prefix}_composer": record["composer"],
                    f"{prefix}_pro_affiliation": record["pro_affiliation"],
                    f"{prefix}_cae_ipi": (
                        "" if record["cae_ipi"] is None else str(record["cae_ipi"])
                    ),
                    f"{prefix}_split": f"{read_split_value(record['split'], split_default):.2f}",
                    f"{prefix}_spotify_link": record.get("spotify_link", ""),
                    f"{prefix}_apple_music_link": record.get("apple_music_link", ""),
                    f"{prefix}_artist_name": record["artist_name"],
                }
            )

    return state_updates


def build_default_form_state() -> dict[str, object]:
    return {
        "album_name": "",
        "producer": "",
        "track_count": 1,
        "show_track_info_import": False,
        "_composer_lookup": {},
        "_show_export_success_dialog": False,
        "_export_success_nonce": 0,
        "track_1_title": "",
        "track_1_bpm": 0,
        "track_1_key": "",
        "track_1_meter": "",
        "track_1_meter_numerator": "",
        "track_1_meter_denominator": "",
        "track_1_instrumentation": [],
        "track_1_has_lyrics": False,
        "track_1_lyrics": "",
        "track_1_vocal_sublist": [],
        "track_1_featured_instrument": None,
        "track_1_multi_key_enabled": False,
        "track_1_multi_key_previous": False,
        "track_1_multi_key_count": 1,
        "track_1_single_key": None,
        "track_1_multi_keys": [],
        "track_1_composer_count": 1,
        f"{composer_prefix(1, 1)}_saved": "",
        f"{composer_prefix(1, 1)}_composer": "",
        f"{composer_prefix(1, 1)}_pro_affiliation": "",
        f"{composer_prefix(1, 1)}_cae_ipi": "",
        f"{composer_prefix(1, 1)}_split": "100.00",
        f"{composer_prefix(1, 1)}_spotify_link": "",
        f"{composer_prefix(1, 1)}_apple_music_link": "",
        f"{composer_prefix(1, 1)}_artist_name": "",
    }


def apply_pending_import() -> None:
    imported_workbook = st.session_state.pop("_pending_track_info_import", None)
    if not imported_workbook:
        return

    keys_to_clear = [
        key
        for key in list(st.session_state.keys())
        if key == "_composer_lookup"
        or key == "track_count"
        or re.match(r"track_\d+_", key)
    ]
    for key in keys_to_clear:
        del st.session_state[key]

    for key, value in build_import_state(imported_workbook).items():
        st.session_state[key] = value


def apply_pending_reset() -> None:
    if not st.session_state.pop("_pending_track_info_reset", False):
        return

    keys_to_clear = [
        key
        for key in list(st.session_state.keys())
        if not str(key).startswith("$$STREAMLIT_INTERNAL_KEY")
    ]
    for key in keys_to_clear:
        del st.session_state[key]

    for key, value in build_default_form_state().items():
        st.session_state[key] = value


def toggle_import_panel() -> None:
    st.session_state["show_track_info_import"] = not st.session_state.get(
        "show_track_info_import",
        False,
    )


def render_import_tool() -> None:
    st.session_state.setdefault("show_track_info_import", False)
    import_message = st.session_state.pop("_track_info_import_message", "")
    import_error = st.session_state.pop("_track_info_import_error", "")

    with st.container(key="track-info-import-toggle"):
        st.button(
            "Need to fix Track Info?",
            key="toggle_track_info_import",
            on_click=toggle_import_panel,
            use_container_width=True,
        )

    if import_message:
        st.success(import_message)
    if import_error:
        st.error(import_error)

    if not st.session_state.get("show_track_info_import", False):
        return

    uploaded_file = st.file_uploader(
        "Upload a previously exported Track Info zip file",
        type=["zip"],
        key="track_info_import_file",
    )
    if uploaded_file is not None:
        st.caption("This exported zip already includes the Track Info Excel file and the Lyrics document when one exists.")
    if st.button(
        "Load Export Into Form",
        key="load_track_info_import",
        disabled=uploaded_file is None,
        use_container_width=True,
    ):
        try:
            imported_workbook, matched_count, unmatched_titles = parse_imported_bundle(
                uploaded_file.getvalue()
            )
            message_parts = [
                f"Loaded {len(imported_workbook['tracks'])} track(s) from the zip bundle."
            ]
            if matched_count:
                message_parts.append(f"Matched lyrics for {matched_count} track(s).")
            if unmatched_titles:
                unmatched_preview = ", ".join(unmatched_titles[:3])
                if len(unmatched_titles) > 3:
                    unmatched_preview += ", ..."
                message_parts.append(
                    f"Could not match lyrics for: {unmatched_preview}"
                )
        except Exception as exc:
            st.session_state["_track_info_import_error"] = str(exc)
            st.rerun()
        else:
            st.session_state["_pending_track_info_import"] = imported_workbook
            st.session_state["_track_info_import_message"] = " ".join(message_parts)
            st.session_state["show_track_info_import"] = False
            st.rerun()


def validation_messages(tracks: list[dict[str, object]]) -> list[str]:
    messages: list[str] = []

    if not tracks:
        messages.append("Track Count: at least one track is required.")
        return messages

    if not tracks[0]["album_name"]:
        messages.append("Album Name is required.")
    if not tracks[0]["producer"]:
        messages.append("WCPM Producer is required.")

    for track in tracks:
        track_number = track["track_number"]
        if not track["track_title"]:
            messages.append(f"Track {track_number}: Title is required.")
        if not track["bpm"]:
            messages.append(f"Track {track_number}: BPM is required.")
        if not track["key"]:
            messages.append(f"Track {track_number}: Key is required.")
        if not track["meter_numerator"] or not track["meter_denominator"]:
            messages.append(f"Track {track_number}: Meter is required.")
        elif not track["meter"]:
            messages.append(
                f"Track {track_number}: Meter must use numbers in a format like 4/4."
            )
        if not track["instrumentation"]:
            messages.append(f"Track {track_number}: Instrumentation is required.")
        if track["has_lyrics"] and not track["lyrics"]:
            messages.append(f"Track {track_number}: Lyrics are required when Has Lyrics? is checked.")
        if (
            includes_vocals(str(track["instrumentation"]).split(", "))
            and not track["vocal_sublist"]
        ):
            messages.append(f"Track {track_number}: Vocal Sub-list is required.")
        if not track["featured_instrument"]:
            messages.append(f"Track {track_number}: Featured Instrument is required.")
        if not track["composers"]:
            messages.append(f"Track {track_number}: Composer information is required.")
        else:
            for composer_index, record in enumerate(track["composers"], start=1):
                if not record["composer"]:
                    messages.append(
                        f"Track {track_number}, Composer {composer_index}: "
                        "Composer is required."
                    )
                if not record["pro_affiliation"]:
                    messages.append(
                        f"Track {track_number}, Composer {composer_index}: "
                        "PRO Affiliation is required."
                    )
                if record["cae_ipi"] is None:
                    messages.append(
                        f"Track {track_number}, Composer {composer_index}: "
                        "CAE/IPI is required."
                    )
                if record["split"] is None:
                    messages.append(
                        f"Track {track_number}, Composer {composer_index}: "
                        "Splits must be a number between 0 and 100."
                    )

            if all(record["split"] is not None for record in track["composers"]):
                split_total = sum(
                    read_split_value(record["split"])
                    for record in track["composers"]
                )
                if abs(split_total - 100.0) > 0.01:
                    messages.append(
                        f"Track {track_number}: composer splits total {split_total:g}%, "
                        "but should equal 100%."
                    )

    return messages


def safe_export_name_part(album_name: str) -> str:
    sanitized = re.sub(r'[<>:"/\\\\|?*]+', " ", str(album_name or "")).strip()
    sanitized = re.sub(r"\s+", " ", sanitized)
    return sanitized or "Track Info"


def safe_filename(album_name: str) -> str:
    return f"{safe_export_name_part(album_name)}_COMPOSER INFO.xlsx"


def safe_lyrics_filename(album_name: str) -> str:
    return f"{safe_export_name_part(album_name)}_LYRICS.docx"


def safe_bundle_filename(album_name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", safe_export_name_part(album_name)).strip("_")
    if not slug:
        slug = "Track_Info"
    timestamp = datetime.now().strftime("%Y%m%d")
    return f"{slug}_Track_Info_Bundle_{timestamp}.zip"


def build_export_bundle(tracks: list[dict[str, object]], album_name: str) -> bytes:
    workbook_file_name = safe_filename(album_name)
    workbook_bytes = build_excel_workbook(tracks)
    lyric_tracks = tracks_with_lyrics(tracks)

    output = BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as export_zip:
        export_zip.writestr(workbook_file_name, workbook_bytes)
        if lyric_tracks:
            export_zip.writestr(
                safe_lyrics_filename(album_name),
                build_lyrics_docx(tracks),
            )
    output.seek(0)
    return output.getvalue()


def parse_imported_bundle(file_bytes: bytes) -> tuple[dict[str, object], int, list[str]]:
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as bundle_zip:
            bundle_members = [
                info.filename
                for info in bundle_zip.infolist()
                if not info.is_dir() and not info.filename.startswith("__MACOSX/")
            ]
            workbook_members = [
                member for member in bundle_members if member.lower().endswith(".xlsx")
            ]
            lyrics_members = [
                member for member in bundle_members if member.lower().endswith(".docx")
            ]

            if not workbook_members:
                raise ValueError("This zip file does not contain a Track Info Excel file.")
            if len(workbook_members) > 1:
                raise ValueError("This zip file contains multiple Excel files. Please keep only one Track Info export in the bundle.")
            if len(lyrics_members) > 1:
                raise ValueError("This zip file contains multiple Word files. Please keep only one Lyrics export in the bundle.")

            imported_workbook = parse_imported_workbook(bundle_zip.read(workbook_members[0]))
            matched_count = 0
            unmatched_titles: list[str] = []

            if lyrics_members:
                imported_lyrics = parse_imported_lyrics_docx(bundle_zip.read(lyrics_members[0]))
                matched_count, unmatched_titles = merge_imported_lyrics(
                    imported_workbook,
                    imported_lyrics,
                )
    except zipfile.BadZipFile as exc:
        raise ValueError("This file is not a valid zip archive.") from exc

    return imported_workbook, matched_count, unmatched_titles


def mark_export_success(file_name: str) -> None:
    st.session_state["_show_export_success_dialog"] = True
    st.session_state["_export_success_file_name"] = file_name
    st.session_state["_export_success_nonce"] = (
        int(st.session_state.get("_export_success_nonce", 0)) + 1
    )


def clear_export_success_dialog() -> None:
    st.session_state["_show_export_success_dialog"] = False
    st.session_state.pop("_export_success_file_name", None)


def reset_form_state() -> None:
    st.session_state["_pending_track_info_reset"] = True


@st.dialog(
    "Track Info Exported",
    width="small",
    icon=":material/check_circle:",
    on_dismiss=clear_export_success_dialog,
)
def render_export_success_dialog(export_success_nonce: int) -> None:
    _ = export_success_nonce
    st.success("Your Track Info zip was successfully exported.")
    export_file_name = compact_text(st.session_state.get("_export_success_file_name", ""))
    if export_file_name:
        st.caption(export_file_name)

    track_count = int(st.session_state.get("track_count", 1))
    tracks = collect_tracks(track_count)
    if tracks_with_lyrics(tracks):
        st.caption("The zip includes both the Track Info Excel file and the Lyrics Word file.")
    else:
        st.caption("The zip includes the Track Info Excel file.")

    if st.button(
        "Close",
        key="close_export_success_dialog",
        use_container_width=True,
    ):
        clear_export_success_dialog()
        st.rerun()


def render_export(track_count: int) -> None:
    tracks = collect_tracks(track_count)
    messages = validation_messages(tracks)

    if messages:
        with st.expander("Required Fields", expanded=False):
            for message in messages:
                st.warning(message)

    album_name = compact_text(st.session_state.get("album_name", ""))
    bundle_bytes = build_export_bundle(tracks, album_name)
    bundle_file_name = safe_bundle_filename(album_name)
    action_cols = st.columns([4, 1])
    with action_cols[0]:
        with st.container(key="track-info-export-info"):
            st.download_button(
                "Export Info",
                data=bundle_bytes,
                file_name=bundle_file_name,
                mime="application/zip",
                key="export_info",
                on_click=mark_export_success,
                args=(bundle_file_name,),
                disabled=bool(messages),
                type="primary",
                use_container_width=True,
            )
    with action_cols[1]:
        with st.container(key="track-info-reset"):
            if st.button(
                "Reset",
                key="reset_track_info",
                use_container_width=True,
            ):
                reset_form_state()
                st.rerun()

    if st.session_state.get("_show_export_success_dialog", False):
        render_export_success_dialog(int(st.session_state.get("_export_success_nonce", 0)))


def main() -> None:
    configure_page()
    apply_pending_reset()
    apply_pending_import()

    render_page_help()
    render_logo_header()
    st.title(APP_TITLE)
    st.markdown(
        (
            f'<p class="intro-copy">{INTRO_TEXT}'
            f'<strong class="intro-emphasis">{INTRO_TEXT_EMPHASIS}</strong>'
            "</p>"
        ),
        unsafe_allow_html=True,
    )

    track_count = render_header_fields()
    instrument_options = get_instrument_options()
    vocal_options = get_vocal_options()
    key_options = get_key_options()

    st.divider()
    for track_number in range(1, track_count + 1):
        render_track_fields(
            track_number,
            instrument_options,
            vocal_options,
            key_options,
        )

    st.divider()
    render_export(track_count)
    st.divider()
    render_import_tool()


if __name__ == "__main__":
    main()
