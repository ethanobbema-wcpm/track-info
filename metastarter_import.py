"""Read legacy metadata as data only; never execute workbook or document content."""
from __future__ import annotations

from io import BytesIO
import math
import re
import unicodedata
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 80 * 1024 * 1024
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def text(value: object) -> str:
    return ' '.join(str(value if value is not None else '').split())


def token(value: object) -> str:
    return re.sub(r'[\W_]', '', unicodedata.normalize('NFKC', text(value)).casefold())


def checked_archive(data: bytes) -> zipfile.ZipFile:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError('Each import file must be 20 MB or smaller.')
    archive = zipfile.ZipFile(BytesIO(data))
    if sum(entry.file_size for entry in archive.infolist()) > MAX_EXPANDED_BYTES:
        archive.close()
        raise ValueError('The expanded import file is too large.')
    return archive


def bpm_value(value: object, title: str, warnings: list[str]) -> int | None:
    if not text(value):
        return None
    try:
        number = float(text(value))
        if not math.isfinite(number) or not 0 < number <= 400:
            raise ValueError
        return round(number)
    except ValueError:
        warnings.append(f'{title}: BPM {text(value)!r} could not be imported. Enter it manually.')
        return None


def empty_track(title: str, number: int) -> dict:
    return dict(track_title=title, track_number=number, album_name='', producer='', bpm=None,
                key='', meter='', instrumentation=[], vocal_sublist=[], featured_instrument='',
                has_lyrics=False, lyrics='', composers=[])


COMPOSER = re.compile(
    r'\s*(?P<name>[^\(]+?)\s*\((?P<pro>[^)]+)\)\s*'
    r'(?P<split>\d+(?:\.\d+)?)\s*%\s*\[(?P<ipi>[^\]]+)\]\s*(?:[,;\n]|$)'
)
URL = re.compile(r'https?://[^\s,;)]+', re.I)


def composers_from_cell(value: object, title: str, warnings: list[str]) -> list[dict]:
    raw = text(value)
    records = []
    for match in COMPOSER.finditer(raw):
        ipi = re.sub(r'[\s-]', '', match['ipi'])
        split = float(match['split'])
        records.append(dict(composer=text(match['name']), pro_affiliation=text(match['pro']),
                            cae_ipi=ipi if ipi.isdigit() else None,
                            split=split if 0 <= split <= 100 else None,
                            spotify_link='', apple_music_link='', dsp_link='', artist_name=''))
    leftover = COMPOSER.sub('', raw).strip(' ,;')
    if leftover:
        warnings.append(f'{title}: composer text needs review: {leftover}')
        records.append(dict(composer=leftover, pro_affiliation='', cae_ipi=None, split=None,
                            spotify_link='', apple_music_link='', dsp_link='', artist_name=''))
    return records


def parse_metastarter(data: bytes) -> tuple[dict, list[str]]:
    with checked_archive(data):
        pass
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True, keep_links=False)
    warnings: list[str] = []
    try:
        selected = None
        for sheet in workbook:
            for row_index, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), 1):
                headers = {token(v): i for i, v in enumerate(row) if text(v)}
                if {'tracktitle', 'cdtitle', 'composer'} <= headers.keys():
                    selected = sheet, row_index, headers, row
                    break
            if selected:
                break
        if not selected:
            raise ValueError('No Metastarter table found. Expected TrackTitle, CDTitle, and Composer columns.')
        sheet, header_row, headers, raw_headers = selected
        if sheet.max_row and sheet.max_row > 20000:
            raise ValueError('The Metastarter has too many rows to import.')
        tracks = []
        seen = set()
        volumes = set()
        albums = set()
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            def get(*names):
                for name in names:
                    index = headers.get(token(name))
                    if index is not None and index < len(row):
                        value = row[index]
                        return '' if text(value).startswith(('#REF!', '#VALUE!', '#N/A', '#DIV/0!')) else value
                return ''
            title = text(get('TrackTitle'))
            filename = text(get('Filename'))
            if not title or title.casefold() == 'double-check with column a' or filename.startswith('Volume_'):
                continue
            # Use the main full mix; cutdowns, stems and alternates are not extra compositions.
            if filename and not re.search(r'_Full(?:\.[^.]+)?$', filename, re.I):
                warnings.append(f'Skipped non-Full mix: {filename}')
                continue
            if token(title) in seen:
                raise ValueError(f'Duplicate Full track title: {title}. Resolve it before importing.')
            seen.add(token(title))
            number_match = re.search(r'_(\d+)_', filename)
            track = empty_track(title, int(number_match[1]) if number_match else len(tracks) + 1)
            track.update(album_name=text(get('CDTitle')), producer=text(get('WCPM Producer', 'Producer')),
                         bpm=bpm_value(get('BPM'), title, warnings), key=text(get('Key')),
                         meter=text(get('Meter')), composers=composers_from_cell(get('Composer'), title, warnings),
                         instrumentation=[text(x) for x in text(get('Instrumentation')).split(',') if text(x)],
                         featured_instrument=text(get('Featured Instrument')))
            if not track['composers']:
                track['composers'] = [dict(composer='', pro_affiliation='', cae_ipi=None, split=None,
                                           spotify_link='', apple_music_link='', dsp_link='', artist_name='')]
            artist = text(get('Artist'))
            # Old Metastarters put DSP links in the unnamed column immediately after Artist.
            dsp = text(get('BWDescription (for DSP Artist Info)', 'DSP Links', 'DSP Link'))
            artist_index = headers.get('artist')
            if not dsp and artist_index is not None and artist_index + 1 < len(row) and not text(raw_headers[artist_index + 1]):
                dsp = text(row[artist_index + 1])
            if len(track['composers']) == 1:
                record = track['composers'][0]
                links = URL.findall(dsp)
                record.update(artist_name=artist,
                              spotify_link=next((u for u in links if 'spotify.com/' in u.casefold()), ''),
                              apple_music_link=next((u for u in links if 'music.apple.com/' in u.casefold()), ''))
            elif artist or dsp:
                warnings.append(f'{title}: multiple composers; assign artist names and DSP links manually. Source: {artist} {dsp}')
            tracks.append(track)
            if text(get('Volume')):
                volumes.add(text(get('Volume')))
            if track['album_name']:
                albums.add(track['album_name'])
            if len(tracks) > 200:
                raise ValueError('The app supports at most 200 tracks.')
        if not tracks:
            raise ValueError('No completed Full tracks were found in the Metastarter.')
        if len(albums) > 1 or len(volumes) > 1:
            raise ValueError('Upload a Metastarter for one album at a time.')
        tracks.sort(key=lambda t: t['track_number'])
        return dict(album_name=next(iter(albums), ''), producer=next((t['producer'] for t in tracks if t['producer']), ''), tracks=tracks), warnings
    finally:
        workbook.close()


def merge_track_info(data: bytes, album: dict, warnings: list[str]) -> None:
    with checked_archive(data) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    by_title = {token(t['track_title']): t for t in album['tracks']}
    matched = set()
    table_found = False
    for table in root.iter(W + 'tbl'):
        headers = None
        for row in table.findall(W + 'tr'):
            cells = [
                '\n'.join(''.join(n.text or '' for n in p.iter(W + 't')) for p in c.findall(W + 'p'))
                for c in row.findall(W + 'tc')]
            names = {token(c): i for i, c in enumerate(cells)}
            if 'tracktitle' in names and any(k in names for k in ('bpm', 'instrumentation', 'key', 'meter')):
                headers = names
                table_found = True
                continue
            if headers is None:
                continue
            def get(name):
                i = headers.get(name)
                return text(cells[i]) if i is not None and i < len(cells) else ''
            title = get('tracktitle')
            if not title:
                continue
            title_key = token(title)
            track = by_title.get(title_key)
            if not track:
                warnings.append(f'Word document track did not match the Metastarter: {title}')
                continue
            if title_key in matched:
                raise ValueError(f'Duplicate track in the Word document: {title}')
            matched.add(title_key)
            for field in ('bpm', 'key', 'meter', 'instrumentation', 'featuredinstrument'):
                value = get(field)
                if not value:
                    continue
                target = 'featured_instrument' if field == 'featuredinstrument' else field
                if field == 'bpm':
                    value = bpm_value(value, title, warnings)
                    if value is None:
                        continue
                elif field == 'instrumentation':
                    value = [text(x) for x in re.split(r'[,;\n]', value) if text(x)]
                if track[target] and track[target] != value:
                    warnings.append(f'{title}: used Word document {target} instead of the Metastarter value.')
                track[target] = value
    if not table_found or not matched:
        raise ValueError('No matching Track Info table found in the Word document. Match tracks by title, not row position.')
    for key, track in by_title.items():
        if key not in matched:
            warnings.append(f"No Word document details found for: {track['track_title']}")


INSTRUMENT_ALIASES = {
    'electricbass': 'Bass - Electric', 'electricguitar': 'Guitar - Electric',
    'drumkit': 'Drum Kit', 'programmeddrums': 'Drums - Programmed',
    'synthkeys': 'Synthesizer', 'synthleads': 'Synthesizer', 'synthpads': 'Synth - Pad',
    'synthbass': 'Bass - Synth', 'wahwahguitar': 'Guitar - Wah Wah',
    'electricpiano': 'Piano - Electric', 'hammondorgan': 'Organ - Hammond',
    'recordscratchsfx': 'SFX (Sound Effects) - Record Scratch',
}


def normalize_options(album: dict, instruments: list[str], keys: list[str], warnings: list[str]) -> None:
    options = {token(v): v for v in instruments}
    key_options = {v.casefold(): v for v in keys}
    mappings = set()
    for track in album['tracks']:
        title = track['track_title']
        selected = []
        for source in track['instrumentation']:
            canonical = options.get(token(source)) or options.get(token(INSTRUMENT_ALIASES.get(token(source), '')))
            if canonical:
                if canonical not in selected:
                    selected.append(canonical)
                if source != canonical:
                    mappings.add(f'{source} → {canonical}')
            else:
                warnings.append(f'{title}: instrument not in the app list: {source}. Select it manually.')
        track['instrumentation'] = selected
        featured = track['featured_instrument']
        canonical = options.get(token(featured)) or options.get(token(INSTRUMENT_ALIASES.get(token(featured), '')))
        track['featured_instrument'] = canonical if canonical in selected else ''
        if featured and not track['featured_instrument']:
            warnings.append(f'{title}: featured instrument needs manual selection: {featured}')
        selected_keys = []
        for source in re.split(r'[,/;]', track['key']):
            source = text(source)
            if not source:
                continue
            source = source.replace('♭', 'b').replace('♯', '#')
            source = re.sub(r'\s*major$', '', source, flags=re.I)
            source = re.sub(r'\s*minor$', 'm', source, flags=re.I)
            canonical = key_options.get(source.casefold())
            if not canonical:
                equivalents = {'c#': 'Db', 'd#': 'Eb', 'f#': 'Gb', 'g#': 'Ab', 'a#': 'Bb'}
                root = source[:-1] if source.endswith('m') else source
                equivalent = equivalents.get(root.casefold(), '')
                if equivalent:
                    equivalent += 'm' if source.endswith('m') else ''
                    canonical = key_options.get(equivalent.casefold())
                    if canonical:
                        warnings.append(f'{title}: imported {source} as its equivalent {canonical} from the app key list.')
            if canonical:
                selected_keys.append(canonical)
            else:
                warnings.append(f'{title}: key not in the app list: {source}. Select it manually.')
        track['key'] = ', '.join(selected_keys)
        if track['meter'] and not re.fullmatch(r'\d+\s*/\s*\d+', track['meter']):
            warnings.append(f"{title}: meter needs manual review: {track['meter']}")
            track['meter'] = ''
    if mappings:
        warnings.append('Instrument names mapped to the app list: ' + '; '.join(sorted(mappings)))


def parse_admin_import(workbook: bytes, document: bytes | None, instruments: list[str], keys: list[str]) -> tuple[dict, list[str]]:
    album, warnings = parse_metastarter(workbook)
    if document is not None:
        merge_track_info(document, album, warnings)
    normalize_options(album, instruments, keys, warnings)
    return album, warnings
